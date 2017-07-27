#Andrew Sager
#7/27/2017
import re
import csv
import requests
from openpyxl import Workbook
from uszipcode import ZipcodeSearchEngine

def get_total_pages(url):
	page = requests.get(url.format(1))
	page_source = page.text #get the source code
	num = int(page_source.split(">Last(")[1].split(")")[0])
	return num


def get_zip(url):
	page = requests.get("http://" + url) #append http:// so we can load the page
	page_source = page.text
	pieces = page_source.split("\n")
	#line contains a zip if it contains "postalCode" and "span":
	pieces = [k for k in pieces if ("postalCode" in k and "span" in k)]
	chars_removed = []
	for item in pieces:
		item = re.sub("\D", "", item) #helpful package that removes all non-digits
		if (len(item) < 5): #if we haven't found a proper zipcode
			continue
		else:
			return item[:5] #will truncate any ZIP+4s to ZIPs
	return "" #didn't find any zipcodes in the given url (unlikely)




pages = []
search_url = 'https://www.greenbook.org/advancedsearchresult?page={0}&rd=V&focusgroupid=1&countryId=3'
num_pages = get_total_pages(search_url)
for i in range(num_pages):
	page = requests.get(search_url.format(i+1)) #iterate through all pages of the focus groups
	page_source = page.text #get the source code
	pieces = page_source.split("\n") #break into lines
	pieces = [k for k in pieces if '/company/' in k] #keep the line if it contains a company url
	#continue breaking strings into pieces based on the source code:
	s = "onclick=\"trackOutboundLink('//"
	pieces = [k.split(s) for k in pieces if s in k]
	pieces = [item for sublist in pieces for item in sublist]
	pieces = list(set([k.split("\\r")[0] for k in pieces if ".org" in k]))
	pieces = [k.split("\'")[0] for k in pieces]
	pages.extend(pieces) #add all of the urls we found on this page


search = ZipcodeSearchEngine()
wb = Workbook()
dest_filename = "All Focus Group Centers.xlsx"
new_worksheet = wb.active
new_worksheet.title = "Sheet 1"
new_worksheet.cell(row=1,column=1).value = "Zipcode"
new_worksheet.cell(row=1,column=2).value = "City"
new_worksheet.cell(row=1,column=3).value = "State"
new_worksheet.cell(row=1,column=4).value = "Population"
new_worksheet.cell(row=1,column=5).value = "Density"
new_worksheet.cell(row=1,column=6).value = "URL"
i = 2 #we've already made headers so start at row two
for item in pages:
	zipcode = get_zip(item)
	zipcode_data = search.by_zipcode(zipcode)
	new_worksheet.cell(row=i,column=1).value = zipcode
	new_worksheet.cell(row=i,column=2).value = zipcode_data["City"]
	new_worksheet.cell(row=i,column=3).value = zipcode_data["State"]
	new_worksheet.cell(row=i,column=4).value = zipcode_data["Population"]
	new_worksheet.cell(row=i,column=5).value = zipcode_data["Density"]
	new_worksheet.cell(row=i,column=6).value = item
	i+=1 #increment row to which we add data
wb.save(filename=dest_filename)
print("Compilation of Greenbook data is complete")
